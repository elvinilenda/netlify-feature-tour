"""Exportação da pauta CEJUSC para DOCX com capturas das páginas de prova.

O documento apresenta fatos extraídos literalmente; não é certificação jurídica.
"""
from __future__ import annotations
import io
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PRINTS={'acordao','calculos','homologacao','deposito_recursal','seguro_garantia','sisbajud','mandados','penhoras','alvaras','deposito_judicial'}
FINANCEIRO=re.compile(r'(R\$\s*[\d.,]+|\b\d{1,2}(?:,\d+)?%\b|\bPOSITIVO\b|\bNEGATIVO\b)',re.I)


def _sombreado(paragrafo,fill):
    props=paragrafo._p.get_or_add_pPr()
    shd=OxmlElement('w:shd');shd.set(qn('w:fill'),fill);props.append(shd)


def _marca_run(run,fill):
    props=run._r.get_or_add_rPr()
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),fill);props.append(shade)


def _descricao(p,texto):
    offset=0
    for m in FINANCEIRO.finditer(texto):
        if m.start()>offset:p.add_run(texto[offset:m.start()])
        r=p.add_run(m.group(0));r.bold=True
        offset=m.end()
    if offset<len(texto):p.add_run(texto[offset:])


def gerar_docx(resultado:dict,caminho_pdf:str,max_prints:int=6)->bytes:
    import fitz
    doc=Document();sect=doc.sections[0]
    sect.page_height=Cm(29.7);sect.page_width=Cm(21)
    sect.top_margin=Cm(2);sect.bottom_margin=Cm(2);sect.left_margin=Cm(2.5);sect.right_margin=Cm(2.4)
    estilo=doc.styles['Normal'];estilo.font.name='Arial';estilo.font.size=Pt(10)
    estilo.paragraph_format.space_after=Pt(5)
    fase=resultado.get('fase_identificada') or 'NÃO DETERMINADA'
    verde=fase.startswith('LIQUIDAÇÃO')
    titulo=f'HORÁRIO NÃO INFORMADO - {resultado.get("numero_cnj") or "PROCESSO NÃO IDENTIFICADO"} - {resultado.get("vara") or "VARA NÃO IDENTIFICADA"} - {fase}'
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(10)
    _sombreado(p,'C8F0C3' if verde else 'D8DDE1')
    r=p.add_run(titulo);r.bold=True;r.font.size=Pt(10)
    for chave,rotulo,cor in [('polo_ativo','Polo ativo:','C8F0C3'),('polo_passivo','Polo passivo:','FFE0B5'),('peritos','Peritos:','E2D1FF')]:
        p=doc.add_paragraph();r=p.add_run(rotulo+' ');r.bold=True;_marca_run(r,cor)
        p.add_run(resultado.get(chave) or 'Não localizado nos documentos acessíveis')
    arquivo=fitz.open(caminho_pdf)
    # Prioriza as páginas para print com maior relevância: acórdão, valor líquido,
    # homologação e atos de constrição; não apenas os primeiros comprovantes.
    prioridade={'acordao':0,'homologacao':1,'calculos':2,'sisbajud':3,'penhoras':4,
                'mandados':5,'deposito_recursal':6,'seguro_garantia':7,'deposito_judicial':8,'alvaras':9}
    candidatos=sorted((e for e in resultado.get('eventos',[]) if e['categoria'] in PRINTS and
                        isinstance(e.get('pagina_pdf'),int) and 1<=e['pagina_pdf']<=len(arquivo)),
                       key=lambda e:(prioridade.get(e['categoria'],10),e['pagina_pdf']))
    paginas_print=set()
    for e in candidatos:
        if len(paginas_print)>=max_prints:break
        paginas_print.add(e['pagina_pdf'])
    capturados=set()
    try:
        for e in resultado.get('eventos',[]):
            p=doc.add_paragraph();p.paragraph_format.space_before=Pt(7)
            r=p.add_run(e['data']+' - '+e['ato'].upper()+': ');r.bold=True
            _descricao(p,e['descricao'])
            p.add_run(f' (p. PDF {e["pagina_pdf"]})')
            cat=e['categoria'];pagina=e['pagina_pdf']
            if cat in PRINTS and pagina in paginas_print and pagina not in capturados:
                origem=arquivo[pagina-1]
                caixas=[fitz.Rect(b['bbox']) for b in origem.get_text('dict').get('blocks',[])
                        if b.get('type') in (0,1) and b.get('bbox')]
                if caixas:
                    caixa=caixas[0]
                    for c in caixas[1:]:caixa|=c
                    # Print focado na área do documento, não nas margens vazias.
                    regiao=fitz.Rect(caixa.x0-18,caixa.y0-18,caixa.x1+18,caixa.y1+18)&origem.rect
                else:
                    regiao=origem.rect
                img=origem.get_pixmap(matrix=fitz.Matrix(1.25,1.25),clip=regiao,alpha=False)
                blob=io.BytesIO(img.tobytes('png'))
                pp=doc.add_paragraph();r=pp.add_run(f'Documento para conferência — página {pagina} do PDF original');r.italic=True;r.font.size=Pt(8)
                doc.add_picture(blob,width=Cm(11.3))
                capturados.add(pagina)
        p=doc.add_paragraph();p.add_run('SITUAÇÃO ATUAL: ').bold=True;p.add_run(resultado.get('situacao_atual') or '')
        p=doc.add_paragraph();t=p.add_run('ANÁLISE PRELIMINAR: confira o inteiro teor das peças, a origem dos valores, os resultados das pesquisas e os documentos originais.');t.italic=True;t.font.size=Pt(9)
        output=io.BytesIO();doc.save(output);return output.getvalue()
    finally:
        arquivo.close()